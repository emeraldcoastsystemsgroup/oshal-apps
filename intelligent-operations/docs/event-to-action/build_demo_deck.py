#!/usr/bin/env python3
"""Build the presentation from measured receipts, keeping planned capabilities explicit.
CHANGE LOG: 3 | maintainer@emeraldcoastsystemsgroup.com | Present verified input lineage, current Studio progress and reference cost from the fresh linked monitoring investigation.
CHANGE LOG: 2 | maintainer@emeraldcoastsystemsgroup.com | Include receipt-backed native 3D and all-phase shared-workspace verification.
CHANGE LOG: 1 | maintainer@emeraldcoastsystemsgroup.com | Produce a reproducible TriageNorth demo deck without inventing incident accuracy or autonomous recovery.
Requires python-pptx==1.0.2; the existing document builder produces Word/PDF/Visio.
"""
import json
from pathlib import Path
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont

ROOT=Path(__file__).resolve().parent
EVIDENCE=ROOT.parents[1]/'deploy/evidence'
NAVY='14233B';INK='20334D';TEAL='007F88';BLUE='2364B0';MUTED='576B83';PALE='F3F6FA';MINT='E8F7F4';WHITE='FFFFFF';AMBER='AE6500'
pdfmetrics.registerFont(TTFont('DemoDeck','/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf'))
pdfmetrics.registerFont(TTFont('DemoDeckBold','/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf'))


def build():
    replay=json.loads((EVIDENCE/'replay.json').read_text());baseline=json.loads((EVIDENCE/'recorded-bgl.json').read_text())
    assert replay['passed'] and replay['synthetic'] and baseline['sourceWasSampled']
    actual=replay['result']['actual'];counts=baseline['counts']
    decision=json.loads((EVIDENCE/'decision-chain.json').read_text());workflow=decision['browser'];graph=json.loads((EVIDENCE/'graph-experience.json').read_text())
    assert workflow['passed'] and graph['passed'] and len(workflow['artifacts'])==3
    deck=Presentation();deck.slide_width=Inches(13.333);deck.slide_height=Inches(7.5)

    def text(slide,x,y,w,h,content,size=20,color=INK,bold=False):
        shape=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));frame=shape.text_frame;frame.word_wrap=True
        frame.margin_left=frame.margin_right=frame.margin_top=frame.margin_bottom=0
        spacing=4 if size>=16 else 2
        def height(font_size):
            lines=0
            for paragraph in content.split('\n'):
                current='';lines+=1
                for word in paragraph.split():
                    candidate=(current+' '+word).strip()
                    if current and pdfmetrics.stringWidth(candidate,'DemoDeckBold' if bold else 'DemoDeck',font_size)>w*72:
                        lines+=1;current=word
                    else:current=candidate
            return lines*font_size*1.2+len(content.split('\n'))*spacing
        minimum=min(12,size)
        while size>minimum and height(size)>h*72:size-=1
        if height(size)>h*72:raise ValueError('Deck text does not fit: '+content[:70])
        for i,line in enumerate(content.split('\n')):
            p=frame.paragraphs[0] if i==0 else frame.add_paragraph();p.text=line;p.font.name='DejaVu Sans';p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=RGBColor.from_string(color);p.space_after=Pt(spacing)
        return shape

    def box(slide,x,y,w,h,title,body,tone=MINT):
        shape=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h));shape.fill.solid();shape.fill.fore_color.rgb=RGBColor.from_string(tone);shape.line.color.rgb=RGBColor.from_string('D6E0EB')
        text(slide,x+.2,y+.15,w-.4,.5,title,19,TEAL,True);text(slide,x+.2,y+.8,w-.4,h-.85,body,16)

    def page(title,subtitle):
        s=deck.slides.add_slide(deck.slide_layouts[6]);s.background.fill.solid();s.background.fill.fore_color.rgb=RGBColor.from_string(WHITE)
        text(s,.5,.2,12.3,.3,'TRIAGENORTH  /  SIGNALS → EVIDENCE → ACTION',12,TEAL,True)
        text(s,.5,.8,12.3,.7,title,31,NAVY,True);text(s,.5,1.6,12.3,.6,subtitle,16,MUTED)
        text(s,.5,7.05,12.3,.25,'21 OCT 2026  ·  Demonstrated behavior and planned work are identified separately  ·  '+str(len(deck.slides)),10,MUTED)
        return s

    s=page('Every incident deserves its evidence.','TriageNorth is an intelligent operations experience for mixed infrastructure.')
    text(s,.7,2.8,11.8,1.3,'Turn a flood of signals into a small set of explainable incidents—with the context needed to act.',29,NAVY,True)
    box(s,.7,4.6,5.8,1.8,'The product','Named sources, noise reduction, dependency context and controlled response.')
    box(s,6.8,4.6,5.8,1.8,'The foundation','An independently installed OSHAL container supplies identity, connections, tickets, workflows and workspaces.',PALE)

    s=page('The problem is the gap between alert and action.','An isolated alarm rarely explains the system that failed.')
    for x,title,body in [(.6,'Too much noise','Repeated symptoms compete with genuinely new issues.'),(4.9,'Vanishing context','Pods disappear, logs roll and the original failure window is lost.'),(9.2,'Unclear authority','Evidence, generated proposals and permission to execute need distinct boundaries.')]:box(s,x,2.7,3.5,3.5,title,body,PALE)

    s=page('One path from infrastructure to an accountable decision.','The application reuses shared platform controls rather than holding broad cluster credentials.')
    labels=[('Sources','Hosts / cloud / containers\nOTLP / CloudWatch'),('Detection','Stored queries\nSuppress repeated symptoms'),('Context','Dependency paths\nThree telemetry banks'),('Evidence','Bounded durable packet\nSource / time / hashes')]
    for i,(title,body) in enumerate(labels):box(s,.55+i*3.2,2.7,2.95,2.45,title,body)
    text(s,.7,5.7,11.9,.8,'Next boundary: canonical investigation workflow → generated deliverables → exact-action approval → recovery verification.',20,NAVY,True)
    text(s,.7,6.5,11.9,.4,'The replay demonstrates filtering through persistent evidence. Full unattended remediation is not claimed.',14,MUTED)

    s=page('A failure crosses the datacenter boundary.','Synthetic scenario, actual native Kubernetes services: datacenter host → cloud database → network path.')
    for i,(name,key) in enumerate([('Occurrences','deliveries'),('Rule matches','matched'),('Symptoms','survivors'),('Incident groups','groups')]):box(s,.6+i*3.2,2.7,2.9,2.0,name,str(actual[key]))
    text(s,.75,5.15,11.8,1.3,f"{actual['filtered']} routine occurrences filtered. {actual['suppressed']} repeated symptoms suppressed. {actual['evidencePackets']} persistent context packets.\nAn independent disk-pressure host stays in a separate incident group.",22,NAVY,True)

    s=page('Keep the first fifteen minutes available.','A → B → C → A, with protected context before safe reuse.')
    for i,title in enumerate(['A  /  active','B  /  retained','C  /  retained']):box(s,.7+i*4.25,2.65,3.8,2.2,title,'Bounded byte and record budget\nGrow within an approved ceiling\nKeep original provenance')
    text(s,.8,5.15,11.7,1.1,'Never erase a bank whose newest receipt is younger than 900 seconds. At the ceiling, retry the whole batch. Pin incident evidence outside the ring.',22,NAVY,True)
    text(s,.8,6.4,11.7,.4,'Three logical data banks reuse the application store; physical database/PVC autoscaling is separate work.',13,MUTED)

    s=page('Real data keeps the claims honest.','Recorded BGL sample: 2,000 HPC log lines; labels withheld from the severity-only query.')
    box(s,.7,2.65,3.8,2.0,'Precision',f"{baseline['precision']*100:.1f}%\n{counts['falsePositive']} false positives",PALE)
    box(s,4.9,2.65,3.8,2.0,'Recall',f"{baseline['recall']*100:.1f}%\n{counts['falseNegative']} labeled alerts missed",PALE)
    box(s,9.1,2.65,3.5,2.0,'Reduction',f"{counts['matched']} matches\n{counts['suppressed']} repeated occurrences suppressed",PALE)
    text(s,.8,5.05,11.7,1.2,'A severity label alone is a poor incident decision. This measured baseline gives us a target for better rules, temporal context and anomaly evaluation.',23,NAVY,True)
    text(s,.8,6.4,11.7,.4,'Loghub research samples remain outside the product bundle. This is not a trained anomaly detector or root-cause accuracy score.',13,MUTED)

    s=page('A demo you can inspect, not just watch.','Each stage shows its sources, actual decisions and retained evidence.')
    for i,(title,body) in enumerate([('Scenario lab','Run the host/cloud scenario and inspect the result.'),('Live context','Inspect all three bank generations and their budgets.'),('Data & topology','Trace filter decisions, symptom keys and dependency paths.')]):box(s,.7+i*4.25,2.7,3.8,3.0,title,body)
    text(s,.8,6.1,11.7,.65,'Login and grants belong to the signed-in user. The investigation receives a preprocessed packet, not unrestricted infrastructure access.',19,NAVY,True)

    s=page('Explore the estate. Follow the real work.','Real monitoring from a synthetic workload; three analysis phases await human report approval.')
    box(s,.7,2.65,3.8,3.0,'Six graph queries','Perspective depth + orbit/pan/zoom.\nRecorded dependencies, impact and placement.\nKeyboard selection and phone layout.')
    box(s,4.95,2.65,3.8,3.0,'Inspect each decision','Named investigate / review / revise bots.\nExact admitted packet and prior reports.\nEvidence citations + proposed diagnostics.\nCurrent position in graphical Studio.')
    box(s,9.2,2.65,3.4,3.0,'Review work and cost','Three shared reports, byte/hash checked.\n39,166 measured tokens.\n$0.050048 reference estimate.\n$0 recorded spend; reference is not billed cost.')
    text(s,.8,6.1,11.7,.65,'Reports and observed correlation do not establish infrastructure recovery. Generated executable remediation and external status synchronization remain delivery gates.',18,MUTED)

    s=page('Connect accountable work to an existing issue.','Development-installed services; full release certification and live Jira account validation remain open.')
    box(s,.7,2.65,3.8,3.0,'Review an exact mapping','Owned workflow revision + phase.\nPersonal named account.\nExplicit Jira / GitLab / ServiceNow state mapping.')
    box(s,4.95,2.65,3.8,3.0,'Confirm one update','Reserve before dispatch.\nUse shared native connector Actions.\nRead state before and after.\nHold uncertain responses.')
    box(s,9.2,2.65,3.4,3.0,'Notify through shared services','SMS / brief phone / Telegram.\nThree approved tests accepted once.\nAll test topics muted.\nOperator confirmed the call.\nSMS blocked: Twilio30034.\nTelegram receipt unverified.')
    text(s,.8,6.1,11.7,.65,'Automatic two-way sync, verified webhooks and infrastructure recovery remain delivery gates. The acceptance runner never repeats an attempted test send.',18,MUTED)

    s=page('The next three weeks focus on function and evidence.','Build on the demonstrated source-to-context path.')
    box(s,.7,2.7,3.8,3.3,'Evaluate','Replay recorded host and cloud metrics.\nCalibrate anomalies on time-separated data.\nKeep failures and false alerts visible.',PALE)
    box(s,4.95,2.7,3.8,3.3,'Complete response','Continuous delegated producers.\nAutomatic evidence-to-workflow dispatch.\nGenerated scripts, exact approvals and failed-recovery checks.',PALE)
    box(s,9.2,2.7,3.4,3.3,'Expand environments','External ticket synchronization.\nRead-only AWS / Azure / GCP discovery.\nEducational cloud environments after the product path is ready.',PALE)
    deck.save(ROOT/'triagenorth-demo.pptx')
    print(json.dumps({'slides':len(deck.slides),'path':str(ROOT/'triagenorth-demo.pptx')}))


if __name__=='__main__':build()
